import asyncio
from dataclasses import replace
from io import BytesIO
import json
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from app.columns import SETTINGS
from app.enrichment import NaicsReference, ProviderResult
from app.job_models import JobMetadata, ScoreResult
from app.main import create_app
from app.persistence import JobRepository, LocalArtifactStore
from app.pipeline import Pipeline


def xlsx(rows, headers=("Business Name", "NAICS", "ZIP"), sheet_name="Accounts"):
    workbook = Workbook()
    workbook.active.title = sheet_name
    workbook.active.append(headers)
    for row in rows:
        workbook.active.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def setup(tmp_path):
    settings = replace(SETTINGS, job_artifact_directory=str(tmp_path), naics_retry_base_seconds=0, worker_poll_interval=0.01)
    repo, store = JobRepository(tmp_path), LocalArtifactStore(tmp_path)
    repo.initialize()
    return settings, repo, store


def submit(client, data, metadata=None, filename="input.xlsx"):
    metadata = metadata or {"sheet_name": "Accounts", "mapping": {"business_name": "A", "naics": "B", "postal_code": "C"}}
    return client.post("/api/v1/jobs", files={"file": (filename, data)}, data={"metadata": json.dumps(metadata)})


class FakeProvider:
    name, version = "fake", "1"

    def __init__(self):
        self.calls = []

    async def lookup(self, identity):
        self.calls.append(identity.business_name)
        return ProviderResult(status="not_found" if identity.business_name == "No match" else "found",
                              naics=None if identity.business_name == "No match" else "541330", record_id="record-1", confidence=0.99)


class FakeScorer:
    name, version, required_fields = "Score A", "1", ("naics",)

    def score(self, features):
        return {number: ScoreResult(value=0.73, status="scored") for number in features}


def test_six_row_acceptance_and_download(setup):
    settings, repo, store = setup
    rows = [["Supplied", "541330", "00123"], ["Enriched", None, "00123"], ["Replaced", "bad", "00123"],
            ["No match", None, "00123"], [None, None, None], ["Enriched", None, "00123"]]
    data = xlsx(rows)
    provider = FakeProvider()
    with TestClient(create_app(settings=settings)) as client:
        schema = client.get("/api/schema").json()
        assert next(field for field in schema["fields"] if field["key"] == "business_name")["required"]
        assert not next(field for field in schema["fields"] if field["key"] == "naics")["required"]
        accepted = submit(client, data)
        assert accepted.status_code == 202
        job_id = accepted.json()["job_id"]
        assert client.get(f"/api/v1/jobs/{job_id}/result").status_code == 409
        job = repo.claim()
        assert repo.claim() is None
        pipeline = Pipeline(repo, store, settings, provider, NaicsReference({"541330"}, "2022-test"), [FakeScorer()])
        asyncio.run(pipeline.process(job))
        status = client.get(f"/api/v1/jobs/{job_id}").json()
        assert status["status"] == "completed_with_issues"
        assert (status["total_rows"], status["processed_rows"], status["scored_rows"], status["needs_review_rows"], status["invalid_rows"]) == (6, 6, 4, 1, 1)
        response = client.get(f"/api/v1/jobs/{job_id}/result")
        assert response.status_code == 200
        result = load_workbook(BytesIO(response.content))
        values = list(result["Accounts"].values)
        assert len(values) - 1 == len(rows)
        assert [list(row[:3]) for row in values[1:]] == rows
        assert [row[3] for row in values[1:]] == list(range(2, 8))
        assert [row[4] for row in values[1:]] == ["scored", "scored", "scored_with_warnings", "needs_review", "invalid", "scored"]
        assert values[3][7] == "bad" and values[3][8] == "541330" and values[3][9] == "third_party"
        page = client.get(f"/api/v1/jobs/{job_id}/rows").json()
        assert page["total"] == 6 and page["offset"] == 0 and page["limit"] == 50
        assert [row["source_row_number"] for row in page["rows"]] == list(range(2, 8))
        assert [row["status"] for row in page["rows"]] == [row[4] for row in values[1:]]
        for row, excel_row in zip(page["rows"], values[1:]):
            assert row["naics"]["input_value"] == excel_row[7]
            assert row["naics"]["final_value"] == excel_row[8]
            assert row["scores"]["Score A"]["value"] == excel_row[13]
            assert row["canonical_account"]["postal_code"] == ("00123" if row["source_row_number"] != 6 else None)
        with repo.connection() as db:
            saved = [json.loads(row[0]) for row in db.execute("SELECT result FROM batch_job_rows WHERE job_id=? ORDER BY source_row_number", (job_id,))]
        assert page["rows"] == saved
        assert len(provider.calls) == 3 and provider.calls.count("Enriched") == 1
        assert "Run Summary" in result.sheetnames
        assert "fake" in json.dumps(status["versions"])


def test_production_placeholder_never_scores_and_optional_naics(setup):
    settings, repo, store = setup
    with TestClient(create_app(settings=settings)) as client:
        accepted = submit(client, xlsx([["  Alpha  "]], headers=["Applicant"]),
                          {"sheet_name": "Accounts", "mapping": {"business_name": "A"}})
        assert accepted.status_code == 202
        asyncio.run(Pipeline(repo, store, settings).process(repo.claim()))
        job = repo.get(accepted.json()["job_id"])
        assert job["status"] == "completed_with_issues" and job["scored_rows"] == 0
        with repo.connection() as db:
            row = json.loads(db.execute("SELECT result FROM batch_job_rows").fetchone()[0])
        assert row["status"] == "needs_review" and "SCORING_NOT_CONFIGURED" in row["issues"]
        assert "NAICS_PROVIDER_NOT_CONFIGURED" in row["issues"]
        assert load_workbook(store.output_path(job["id"])).active.cell(2, 1).value == "  Alpha  "


@pytest.mark.parametrize("metadata", [
    {"sheet_name": "Accounts", "mapping": {"naics": "B"}},
    {"sheet_name": "Accounts", "mapping": {"business_name": "Z"}},
    {"sheet_name": "Accounts", "mapping": {"business_name": "A", "naics": "A"}},
    {"sheet_name": "Accounts", "mapping": {"business_name": "A", "unknown": "B"}},
    {"sheet_name": "Missing", "mapping": {"business_name": "A"}},
])
def test_bad_submissions_are_not_accepted(setup, metadata):
    settings, repo, store = setup
    with TestClient(create_app(settings=settings)) as client:
        response = submit(client, xlsx([["Alpha", "541330", "00123"]]), metadata)
        assert response.status_code == 422 and "error" in response.json()
        assert repo.claim() is None


def test_api_failure_contracts_and_readiness(setup):
    settings, repo, store = setup
    with TestClient(create_app(settings=settings)) as client:
        assert client.get("/api/ready").status_code == 503
        repo.heartbeat()
        assert client.get("/api/ready").status_code == 200
        assert client.get(f"/api/v1/jobs/{uuid4()}").status_code == 404
        assert client.get(f"/api/v1/jobs/{uuid4()}/result").status_code == 404
        assert client.get("/api/v1/jobs/not-an-id").json()["error"]["code"] == "INVALID_REQUEST"
        assert client.post("/api/v1/jobs").status_code == 422
        assert client.post("/api/v1/jobs", files={"file": ("test.xlsx", b"x")}, data={"metadata": "{"}).status_code == 422
        assert submit(client, b"bad").json()["error"]["code"] == "INVALID_WORKBOOK"
        assert submit(client, b"bad", filename="test.xls").status_code == 415
        job_id = submit(client, xlsx([["Alpha", "541330", "12345"]]), filename="../../original.xlsx").json()["job_id"]
        assert repo.get(job_id)["original_filename"] == "original.xlsx"
        repo.update(job_id, status="completed")
        assert client.get(f"/api/v1/jobs/{job_id}/result").status_code == 503


def test_complete_job_and_restart_checkpoint(setup):
    settings, repo, store = setup
    provider = FakeProvider()
    with TestClient(create_app(settings=settings)) as client:
        job_id = submit(client, xlsx([["Alpha", None, "12345"]])).json()["job_id"]
        pipeline = Pipeline(repo, store, settings, provider, NaicsReference({"541330"}, "test"), [FakeScorer()])
        asyncio.run(pipeline.process(repo.claim()))
        assert repo.get(job_id)["status"] == "completed"
        repo.update(job_id, status="running")
        repo.recover(3)
        asyncio.run(pipeline.process(repo.claim()))
        assert provider.calls == ["Alpha"]
        assert repo.get(job_id)["status"] == "completed"


def test_failed_job_never_downloads_and_retries_are_bounded(setup):
    settings, repo, store = setup
    with TestClient(create_app(settings=settings)) as client:
        job_id = submit(client, xlsx([["Alpha", None, "12345"]])).json()["job_id"]
        (store.root / job_id / "input.xlsx").unlink()
        pipeline = Pipeline(repo, store, settings)
        for attempt in range(settings.worker_max_attempts):
            asyncio.run(pipeline.process(repo.claim()))
        assert repo.get(job_id)["status"] == "failed"
        assert client.get(f"/api/v1/jobs/{job_id}/result").status_code == 409
        # Corrupt immutable input is a workbook-level fatal error, not an automatic retry.
        job_id = submit(client, xlsx([["Alpha", None, "12345"]])).json()["job_id"]
        store.save_input(job_id, b"corrupt")
        asyncio.run(pipeline.process(repo.claim()))
        assert repo.get(job_id)["status"] == "failed"


def test_restart_does_not_mix_processing_versions(setup):
    settings, repo, store = setup
    with TestClient(create_app(settings=settings)) as client:
        job_id = submit(client, xlsx([["Alpha", "541330", "12345"]])).json()["job_id"]
        asyncio.run(Pipeline(repo, store, settings).process(repo.claim()))
        original_versions = repo.get(job_id)["versions"]
        repo.update(job_id, status="running")
        repo.recover(3)
        changed = replace(settings, pipeline_version="changed")
        asyncio.run(Pipeline(repo, store, changed).process(repo.claim()))
        job = repo.get(job_id)
        assert job["status"] == "failed" and job["error_code"] == "PIPELINE_VERSION_CHANGED"
        assert job["versions"] == original_versions
        assert client.get(f"/api/v1/jobs/{job_id}/result").status_code == 409


def test_second_job_uses_cross_job_cache(setup):
    settings, repo, store = setup
    provider = FakeProvider()
    with TestClient(create_app(settings=settings)) as client:
        pipeline = Pipeline(repo, store, settings, provider, NaicsReference({"541330"}, "test"), [FakeScorer()])
        for _ in range(2):
            submit(client, xlsx([["Alpha", None, "12345"]]))
            asyncio.run(pipeline.process(repo.claim()))
        assert len(provider.calls) == 1


def test_five_thousand_rows_remain_aligned_with_distinct_scores(setup):
    settings, repo, store = setup
    rows = [[f" Business {number} ", "541330", "00123"] for number in range(5000)]

    class RowScore(FakeScorer):
        def score(self, features):
            return {number: ScoreResult(value=number, status="scored") for number in features}

    with TestClient(create_app(settings=settings)) as client:
        job_id = submit(client, xlsx(rows)).json()["job_id"]
        asyncio.run(Pipeline(repo, store, settings, naics_reference=NaicsReference({"541330"}, "test"), scorers=[RowScore()]).process(repo.claim()))
        job = repo.get(job_id)
        assert job["status"] == "completed" and job["scored_rows"] == 5000
        output = load_workbook(store.output_path(job_id), read_only=True)
        values = list(output.active.values)
        headers = list(values[0])
        score_position = headers.index("Score A")
        assert len(values) == 5001
        for number, row in enumerate(values[1:], 2):
            assert list(row[:3]) == rows[number - 2]
            assert row[3] == row[score_position] == number
        assert all(duration >= 0 for duration in job["metrics"]["stage_seconds"].values())
        output.close()
