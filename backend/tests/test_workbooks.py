from dataclasses import replace
from datetime import datetime
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
import pytest

from app.job_models import ApplicationError, JobMetadata, NaicsResolution, RowResult, ScoreResult
from app.workbooks import canonical_rows, input_naics, read_workbook, render_workbook, validate_mapping
from test_jobs import setup, xlsx


def rewrite_archive(data, replacements):
    buffer = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(buffer, "w") as target:
        for name in source.namelist():
            target.writestr(name, replacements.get(name, source.read(name)))
        for name, content in replacements.items():
            if name not in source.namelist():
                target.writestr(name, content)
    return buffer.getvalue()


@pytest.mark.parametrize("overrides,code", [
    ({"max_file_bytes": 1}, "UPLOAD_TOO_LARGE"),
    ({"max_expanded_bytes": 1}, "WORKBOOK_TOO_LARGE"),
    ({"max_workbook_cells": 1}, "WORKBOOK_TOO_LARGE"),
    ({"max_rows": 1}, "WORKBOOK_TOO_LARGE"),
])
def test_limits(setup, overrides, code):
    settings, repo, store = setup
    with pytest.raises(ApplicationError) as error:
        read_workbook(xlsx([["One"], ["Two"]]), "input.xlsx", "Accounts", replace(settings, **overrides))
    assert error.value.code == code


@pytest.mark.parametrize("rows,headers,message", [([], ["Name"], "no data"), ([["One"]], [None], "headers")])
def test_invalid_tables(setup, rows, headers, message):
    with pytest.raises(ApplicationError, match=message):
        read_workbook(xlsx(rows, headers=headers), "input.xlsx", "Accounts", setup[0])


def test_macros_and_invalid_xml_rejected(setup):
    data = xlsx([["One"]])
    with pytest.raises(ApplicationError, match="Macro"):
        read_workbook(rewrite_archive(data, {"xl/vbaProject.bin": b"macros"}), "input.xlsx", "Accounts", setup[0])
    with pytest.raises(ApplicationError, match="read"):
        read_workbook(rewrite_archive(data, {"xl/workbook.xml": b"malformed"}), "input.xlsx", "Accounts", setup[0])
    with ZipFile(BytesIO(data)) as archive:
        truncated_sheet = archive.read("xl/worksheets/sheet1.xml").replace(b"</worksheet>", b"")
    with pytest.raises(ApplicationError, match="worksheet"):
        read_workbook(rewrite_archive(data, {"xl/worksheets/sheet1.xml": truncated_sheet}), "input.xlsx", "Accounts", setup[0])


def test_formulas_errors_dates_and_literal_cells(setup):
    data = xlsx([["Alpha", "=1+1", datetime(2025, 1, 1)], ["Beta", "=2+2", "#N/A"], ["=literal", None, "Text"]])
    with ZipFile(BytesIO(data)) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml").replace(b"<f>1+1</f><v />", b"<f>1+1</f><v>2</v>")
    data = rewrite_archive(data, {"xl/worksheets/sheet1.xml": sheet})
    workbook = read_workbook(data, "input.xlsx", "Accounts", setup[0])
    assert workbook.rows[2][1] == 2
    assert "FORMULA_WITHOUT_CACHED_VALUE" in workbook.cell_issues[3]
    assert "EXCEL_CELL_ERROR" in workbook.cell_issues[3]
    metadata = JobMetadata(sheet_name="Accounts", mapping={"business_name": "A", "naics": "B"})
    accounts = canonical_rows(workbook, metadata)
    assert accounts[1].naics is None and accounts[2].business_name is None
    assert input_naics(workbook, JobMetadata(sheet_name="Accounts", mapping={"business_name": "A", "naics": "C"}), 2) == "2025-01-01T00:00:00"
    results = [RowResult(source_row_number=number, status="needs_review", issues=[], naics=NaicsResolution(status="none")) for number in workbook.rows]
    output = load_workbook(BytesIO(render_workbook(workbook, results, {"Value": "=1+1"})))
    assert output.active.cell(3, 2).value == "=2+2" and output.active.cell(3, 2).data_type == "s"
    assert output["Run Summary"].cell(1, 2).data_type == "s"


def test_output_alignment_duplicates_collisions_and_summary_name(setup):
    workbook = read_workbook(xlsx([["Alpha", None, "value"], [], ["Beta", "X", "Y"]],
                                 headers=["Processing Status", None, "Processing Status"], sheet_name="Run Summary"),
                             "input.xlsx", "Run Summary", setup[0])
    metadata = JobMetadata(sheet_name="Run Summary", mapping={"business_name": "A"}, confirmedFields=["business_name"])
    validate_mapping(workbook, metadata, setup[0])
    accounts = canonical_rows(workbook, metadata)
    assert [row.source_row_number for row in accounts] == [2, 3, 4]
    results = [RowResult(source_row_number=number, status="needs_review", issues=[], naics=NaicsResolution(status="none"),
                         scores={"A": ScoreResult(value=1, status="scored")} if number == 2 else {}) for number in workbook.rows]
    output = load_workbook(BytesIO(render_workbook(workbook, results, {})))
    assert output.active.cell(1, 5).value == "Processing Status (2)"
    assert output.active.cell(3, 15).value == "not_scored"
    assert output.active.cell(1, 2).value is None and "Run Summary (2)" in output.sheetnames
    with pytest.raises(ApplicationError, match="match"):
        render_workbook(workbook, list(reversed(results)), {})


def test_trailing_style_only_cells_are_not_source_columns(setup):
    book = Workbook()
    book.active.title = "Accounts"
    book.active.append(["Business Name"])
    book.active.append(["Alpha"])
    book.active.cell(2, 3).font = Font(bold=True)
    book.active.cell(20, 1).font = Font(bold=True)
    buffer = BytesIO()
    book.save(buffer)
    workbook = read_workbook(buffer.getvalue(), "input.xlsx", "Accounts", replace(setup[0], max_rows=1))
    assert workbook.headers == ["Business Name"] and workbook.rows == {2: ["Alpha"]}
