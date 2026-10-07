"""Tabular XLSX interchange. No Excel business logic or formula evaluation."""
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from .accounts import ACCOUNT_FIELDS, CanonicalAccount, normalize
from .job_models import ApplicationError, JobMetadata, RowResult
from .models import Column, ProcessRequest
from .processing import validate_import


@dataclass
class SourceWorkbook:
    sheet_name: str
    headers: list
    columns: list[Column]
    rows: dict[int, list]
    cell_issues: dict[int, list[str]]
    unusable_cells: dict[int, set[int]]


def read_workbook(data: bytes, filename: str, sheet_name: str, settings) -> SourceWorkbook:
    if not filename.lower().endswith(".xlsx"):
        raise ApplicationError("UNSUPPORTED_FILE_TYPE", "Only .xlsx workbooks are supported.", 415)
    if len(data) > settings.max_file_bytes:
        raise ApplicationError("UPLOAD_TOO_LARGE", "The workbook exceeds the file size limit.", 413)
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > settings.max_expanded_bytes:
                raise ApplicationError("WORKBOOK_TOO_LARGE", "The expanded workbook exceeds the size limit.", 413)
            if any(item.filename.endswith("vbaProject.bin") for item in archive.infolist()):
                raise ApplicationError("INVALID_WORKBOOK", "Macro-enabled workbooks are unsupported.")
        source = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except ApplicationError:
        raise
    except Exception as error:
        raise ApplicationError("INVALID_WORKBOOK", "The workbook could not be read.") from error
    try:
        if sheet_name not in source.sheetnames:
            raise ApplicationError("INVALID_WORKBOOK", "The selected worksheet does not exist.")
        sheet, values = source[sheet_name], cached[sheet_name]
        # Reset untrusted worksheet dimensions; stream actual cells under a hard cell budget.
        sheet.reset_dimensions()
        values.reset_dimensions()
        rows, problems, headers, unusable = {}, {}, [], {}
        last_source_row, cell_count, width = 1, 0, 0
        for number, (cells, saved) in enumerate(zip(sheet.iter_rows(), values.iter_rows()), 1):
            cell_count += len(cells)
            if cell_count > settings.max_workbook_cells:
                raise ApplicationError("WORKBOOK_TOO_LARGE", "The worksheet exceeds the row or cell limit.", 413)
            raw, issues, blocked = [], [], set()
            for index, (cell, cached_cell) in enumerate(zip(cells, saved)):
                value = cell.value
                if cell.data_type == "f":
                    if cached_cell.value is None:
                        issues.append("FORMULA_WITHOUT_CACHED_VALUE")
                        blocked.add(index)
                    else:
                        value = cached_cell.value
                if cell.data_type == "e":
                    issues.append("EXCEL_CELL_ERROR")
                    blocked.add(index)
                raw.append(value)
            while raw and raw[-1] is None:
                raw.pop()
            width = max(width, len(raw))
            if number == 1:
                headers = raw
            elif raw:
                rows[number] = raw
                problems[number] = list(dict.fromkeys(issues))
                unusable[number] = blocked
                last_source_row = number
                if last_source_row > settings.max_rows + 1:
                    raise ApplicationError("WORKBOOK_TOO_LARGE", "The worksheet exceeds the row limit.", 413)
        if not any(str(value).strip() for value in headers if value is not None):
            raise ApplicationError("INVALID_WORKBOOK", "Row 1 must contain headers.")
        if last_source_row == 1:
            raise ApplicationError("INVALID_WORKBOOK", "The selected worksheet has no data rows.")
        headers += [None] * (width - len(headers))
        rows = {number: rows.get(number, []) + [None] * (width - len(rows.get(number, [])))
                for number in range(2, last_source_row + 1)}
        columns = [Column(id=get_column_letter(index + 1), label=str(value if value is not None else "").strip())
                   for index, value in enumerate(headers)]
        return SourceWorkbook(sheet_name, headers, columns, rows, problems, unusable)
    except ApplicationError:
        raise
    except Exception as error:
        raise ApplicationError("INVALID_WORKBOOK", "The selected worksheet could not be read.") from error
    finally:
        source.close()
        cached.close()


def validate_mapping(workbook: SourceWorkbook, metadata: JobMetadata, settings) -> None:
    request = ProcessRequest(columns=workbook.columns, rows=[{}], mapping=metadata.mapping,
                             confirmedFields=metadata.confirmedFields)
    try:
        validate_import(request, ACCOUNT_FIELDS, settings)
    except ValueError as error:
        raise ApplicationError("INVALID_COLUMN_MAPPING", str(error)) from error


def canonical_rows(workbook: SourceWorkbook, metadata: JobMetadata) -> list[CanonicalAccount]:
    positions = {column.id: index for index, column in enumerate(workbook.columns)}
    accounts = []
    for number, raw in workbook.rows.items():
        values = {}
        for field in ACCOUNT_FIELDS:
            column = metadata.mapping.get(field.key)
            value = raw[positions[column]] if column is not None else None
            # Uncached formulas/errors are retained in raw output but never become features.
            if column is not None and positions[column] in workbook.unusable_cells.get(number, set()):
                value = None
            values[field.key] = normalize(value, field.key)
        accounts.append(CanonicalAccount(source_row_number=number, **values))
    return accounts


def input_naics(workbook, metadata, number):
    column = metadata.mapping.get("naics")
    if column is None:
        return None
    index = next(index for index, item in enumerate(workbook.columns) if item.id == column)
    value = workbook.rows[number][index]
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def render_workbook(workbook: SourceWorkbook, results: list[RowResult], summary: dict) -> bytes:
    if [row.source_row_number for row in results] != list(workbook.rows):
        raise ApplicationError("RESULT_ALIGNMENT_ERROR", "Results do not match the submitted rows.", 500)
    output = Workbook()
    sheet = output.active
    sheet.title = workbook.sheet_name
    scorer_names = list(dict.fromkeys(name for row in results for name in row.scores))
    extra = ["Source Row Number", "Processing Status", "Issue Codes", "Issues", "NAICS Submitted", "NAICS Final",
             "NAICS Source", "NAICS Status", "NAICS Provider", "NAICS Retrieved At"]
    extra += [label for name in scorer_names for label in (name, f"{name} Status")]
    used = {str(header) for header in workbook.headers}
    unique = []
    for label in extra:
        candidate = label
        suffix = 2
        while candidate in used:
            candidate = f"{label} ({suffix})"
            suffix += 1
        used.add(candidate)
        unique.append(candidate)
    sheet.append(workbook.headers + unique)
    for result in results:
        naics = result.naics
        values = [result.source_row_number, result.status, ";".join(result.issues),
                  " ".join(issue.replace("_", " ").capitalize() + "." for issue in result.issues),
                  naics.input_value, naics.final_value, naics.source, naics.status, naics.provider,
                  naics.retrieved_at.isoformat() if naics.retrieved_at else None]
        for name in scorer_names:
            score = result.scores.get(name)
            values += [score.value if score else None, score.status if score else "not_scored"]
        sheet.append(workbook.rows[result.source_row_number] + values)
    # Treat all supplied strings literally, including formulas, to avoid evaluation on opening.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    summary_sheet = output.create_sheet("Run Summary" if workbook.sheet_name != "Run Summary" else "Run Summary (2)")
    for key, value in summary.items():
        summary_sheet.append([key, str(value)])
        summary_sheet.cell(summary_sheet.max_row, 2).data_type = "s"
    buffer = BytesIO()
    output.save(buffer)
    output.close()
    return buffer.getvalue()
